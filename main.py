import openpyxl
from openpyxl.utils import get_column_letter
from pathlib import Path
import sys
import csv
from copy import copy
from datetime import datetime

def copy_cell_format(source_cell, target_cell):
    """Copy formatting from source cell to target cell"""
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)

def extract_field_name(cell_value):
    """Extract field name from <fieldname> format"""
    if not cell_value:
        return None
    value = str(cell_value).strip()
    if value.startswith('<') and value.endswith('>'):
        return value[1:-1]
    return None

def parse_template_structure(ws):
    """Parse template to find header field mappings and data table structure"""
    header_mappings = {}
    data_header_row = None
    data_mapping_row = None
    
    for row_idx in range(1, ws.max_row + 1):
        row_values = []
        has_brackets = False
        
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row_idx, col_idx).value
            if cell_value:
                value_str = str(cell_value).strip()
                row_values.append(value_str)
                
                field_name = extract_field_name(value_str)
                if field_name:
                    has_brackets = True
                    cell_ref = f"{get_column_letter(col_idx)}{row_idx}"
                    header_mappings[cell_ref] = field_name
        
        if has_brackets and len([v for v in row_values if v.startswith('<') and v.endswith('>')]) >= 3:
            if data_mapping_row is None:
                data_mapping_row = row_idx
                if row_idx > 1:
                    prev_row_values = []
                    for col_idx in range(1, ws.max_column + 1):
                        val = ws.cell(row_idx - 1, col_idx).value
                        if val and not str(val).strip().startswith('<'):
                            prev_row_values.append(val)
                    if prev_row_values:
                        data_header_row = row_idx - 1
    
    return header_mappings, data_header_row, data_mapping_row

def convert_value(value_str):
    """Try to convert string to appropriate type"""
    if not value_str:
        return None
    
    try:
        if '.' in value_str:
            return float(value_str)
        else:
            return int(value_str)
    except (ValueError, AttributeError):
        pass
    
    for fmt in ['%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%m/%d/%Y', '%d/%m/%Y']:
        try:
            return datetime.strptime(value_str, fmt)
        except (ValueError, AttributeError):
            pass
    
    return value_str

def split_csv_by_booking(input_file, output_folder, ctv_template_file, tac_template_file, booking_column):
    """Split CSV file into separate Excel files per booking"""
    output_path = Path(output_folder)
    output_path.mkdir(exist_ok=True)
    
    print(f"Reading {Path(input_file).name}...")
    
    with open(input_file, 'r', encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        all_data = list(reader)
    
    if len(all_data) > 3:
        print("Skipping first 3 rows (old header format)")
        csv_headers = all_data[3]
        while csv_headers and not csv_headers[-1].strip():
            csv_headers.pop()
        print(f"Found {len(csv_headers)} column headers")
        
        all_rows = []
        for row_data in all_data[4:]:
            if any(cell.strip() for cell in row_data):
                row_dict = {}
                for i, header in enumerate(csv_headers):
                    if i < len(row_data):
                        row_dict[header] = row_data[i].strip()
                    else:
                        row_dict[header] = ''
                all_rows.append(row_dict)
    
    print(f"Found {len(all_rows)} total rows")
    print(f"\nCSV Columns ({len(csv_headers)}):")
    for idx, header in enumerate(csv_headers, 1):
        print(f"  {idx}. {header}")
    
    if all_rows:
        print("\nFirst row sample:")
        for header in csv_headers[:8]:
            value = all_rows[0].get(header, '')
            if len(value) > 50:
                value = value[:47] + "..."
            print(f"  {header}: {value}")
        if len(csv_headers) > 8:
            print(f"  ... and {len(csv_headers) - 8} more columns")
    print()
    
    # Load templates
    ctv_template_wb = None
    ctv_template_ws = None
    tac_template_wb = None
    tac_template_ws = None
    
    if ctv_template_file and Path(ctv_template_file).exists():
        print(f"Loading Crossings TV template: {Path(ctv_template_file).name}")
        ctv_template_wb = openpyxl.load_workbook(ctv_template_file)
        ctv_template_ws = ctv_template_wb.active
    
    if tac_template_file and Path(tac_template_file).exists():
        print(f"Loading The Asian Channel template: {Path(tac_template_file).name}")
        tac_template_wb = openpyxl.load_workbook(tac_template_file)
        tac_template_ws = tac_template_wb.active
    
    if not ctv_template_ws and not tac_template_ws:
        print("Warning: No templates found\n")
    
    # Group rows by booking
    print("Analyzing bookings...")
    
    bookings_data = {}
    invalid_bookings = []
    
    for row in all_rows:
        booking = row.get(booking_column, '').strip()
        if booking:
            is_valid = True
            
            if booking.replace(',', '').replace('.', '').replace(' ', '').isdigit():
                is_valid = False
            
            footer_keywords = ['textbox', 'total', 'sum', 'count', 'page', 'report']
            if any(keyword in booking.lower() for keyword in footer_keywords):
                is_valid = False
            
            if booking and not booking[0].isalpha():
                is_valid = False
            
            if is_valid:
                if booking not in bookings_data:
                    bookings_data[booking] = []
                bookings_data[booking].append(row)
            else:
                if booking not in invalid_bookings:
                    invalid_bookings.append(booking)
    
    if invalid_bookings:
        print(f"\nIgnored {len(invalid_bookings)} invalid booking entries:")
        for inv in invalid_bookings[:5]:
            print(f"  - {inv}")
        if len(invalid_bookings) > 5:
            print(f"  ... and {len(invalid_bookings) - 5} more")
    
    print(f"\nFound {len(bookings_data)} unique bookings:")
    for booking, rows in sorted(bookings_data.items()):
        print(f"  {booking}: {len(rows)} spots")
    
    # Create files
    print("\nCreating split files...")
    
    files_with_unplaced = []  # Track files that have unplaced spots
    
    for booking_num, rows in bookings_data.items():
        markets_in_booking = set()
        has_unplaced = False
        
        for row in rows:
            market = row.get('nome2', '').strip()
            if market:
                markets_in_booking.add(market)
            
            # Check for unplaced spots in dateschedule or airtimep
            dateschedule = row.get('dateschedule', '').strip().lower()
            airtime = row.get('airtimep', '').strip().lower()
            if 'unplaced' in dateschedule or 'unplaced' in airtime:
                has_unplaced = True
        
        has_dallas = 'DALLAS' in markets_in_booking
        
        # Sort rows by market, then date, then airtime
        def sort_key(row):
            market = row.get('nome2', '')
            date_str = row.get('dateschedule', '')
            time_str = row.get('airtimep', '')
            
            # Convert date string to comparable format
            # Put "Unplaced" dates at the end
            if 'unplaced' in date_str.lower():
                date_obj = datetime.max  # Sort unplaced to the end
            else:
                try:
                    date_obj = datetime.strptime(date_str, '%m/%d/%Y')
                except:
                    date_obj = datetime.min
            
            return (market, date_obj, time_str)
        
        rows = sorted(rows, key=sort_key)
        
        # Debug: Check for duplicate consecutive rows (shouldn't happen)
        # Remove any accidental duplicates
        unique_rows = []
        prev_row = None
        for row in rows:
            if row != prev_row:
                unique_rows.append(row)
            prev_row = row
        
        if len(unique_rows) != len(rows):
            print(f"  Note: Removed {len(rows) - len(unique_rows)} duplicate rows from {booking_num}")
        
        rows = unique_rows
        
        if has_dallas and tac_template_ws:
            template_ws = tac_template_ws
            template_name = "The Asian Channel"
        elif not has_dallas and ctv_template_ws:
            template_ws = ctv_template_ws
            template_name = "Crossings TV"
        else:
            template_ws = ctv_template_ws if ctv_template_ws else tac_template_ws
            template_name = "Available"
        
        if template_ws:
            header_mappings, data_header_row, data_mapping_row = parse_template_structure(template_ws)
            
            data_column_map = {}
            if data_mapping_row:
                for col_idx in range(1, template_ws.max_column + 1):
                    cell_value = template_ws.cell(data_mapping_row, col_idx).value
                    field_name = extract_field_name(cell_value)
                    if field_name and field_name in csv_headers:
                        data_column_map[field_name] = col_idx
            
            new_wb = openpyxl.Workbook()
            new_ws = new_wb.active
            new_ws.title = template_ws.title
            
            if data_mapping_row:
                for row_idx in range(1, data_mapping_row + 1):
                    for col_idx in range(1, template_ws.max_column + 1):
                        source_cell = template_ws.cell(row_idx, col_idx)
                        target_cell = new_ws.cell(row_idx, col_idx)
                        target_cell.value = source_cell.value
                        copy_cell_format(source_cell, target_cell)
                
                for col_idx in range(1, template_ws.max_column + 1):
                    col_letter = get_column_letter(col_idx)
                    if col_letter in template_ws.column_dimensions:
                        new_ws.column_dimensions[col_letter].width = template_ws.column_dimensions[col_letter].width
                
                for row_idx in range(1, data_mapping_row + 1):
                    if row_idx in template_ws.row_dimensions:
                        new_ws.row_dimensions[row_idx].height = template_ws.row_dimensions[row_idx].height
                
                first_row = rows[0]
                for cell_ref, field_name in header_mappings.items():
                    if field_name in csv_headers:
                        value = first_row.get(field_name, '')
                        
                        # Special handling for IMPORTO2 (Spot Rate)
                        if field_name == 'IMPORTO2':
                            if not value or value.strip() == '':
                                converted = 0
                            else:
                                converted = convert_value(value)
                        else:
                            converted = convert_value(value) if value else None
                        
                        # Convert duration 119 to 120
                        if field_name == 'duration3' and converted == 119:
                            converted = 120
                        
                        target_cell = new_ws[cell_ref]
                        target_cell.value = converted
                
                # Update calculated fields
                new_ws['B6'].value = len(rows)
                
                if 'IMPORTO2' in data_column_map:
                    total = 0
                    for row in rows:
                        val = row.get('IMPORTO2', '0')
                        try:
                            total += float(val) if val else 0
                        except:
                            pass
                    new_ws['B7'].value = total
                
                # CRITICAL: Explicitly clear row 10 entirely to ensure no data there
                for col_idx in range(1, new_ws.max_column + 1):
                    new_ws.cell(data_mapping_row, col_idx).value = None
                
                # NOW write data rows starting at row 11 (and ONLY row 11+)
                current_row = data_mapping_row + 1  # Start at row 11, NOT row 10
                
                new_ws['B6'].value = len(rows)
                
                if 'IMPORTO2' in data_column_map:
                    total = 0
                    for row in rows:
                        val = row.get('IMPORTO2', '0')
                        try:
                            total += float(val) if val else 0
                        except:
                            pass
                    new_ws['B7'].value = total
                    
                    importo_col = data_column_map['IMPORTO2']
                    new_ws.cell(data_mapping_row, importo_col).number_format = '$#,##0.00'
                
                data_start_row = data_mapping_row + 1
                for row_num, data_row in enumerate(rows, start=data_start_row):
                    for csv_field, excel_col in data_column_map.items():
                        value = data_row.get(csv_field, '')
                        converted_value = convert_value(value) if value else None
                        
                        if csv_field == 'IMPORTO2' and (converted_value is None or converted_value == ''):
                            converted_value = 0
                        
                        # Convert duration 119 to 120
                        if csv_field == 'duration3' and converted_value == 119:
                            converted_value = 120
                        
                        target_cell = new_ws.cell(row_num, excel_col)
                        target_cell.value = converted_value
                        
                        source_format_cell = template_ws.cell(data_mapping_row, excel_col)
                        copy_cell_format(source_format_cell, target_cell)
                        
                        if csv_field == 'IMPORTO2':
                            target_cell.number_format = '$#,##0.00'
                        
                        if csv_field == 'dateschedule':
                            target_cell.number_format = 'M/D/YY'
                
                for col_idx in range(1, new_ws.max_column + 1):
                    col_letter = get_column_letter(col_idx)
                    max_length = 0
                    for row_idx in range(1, new_ws.max_row + 1):
                        cell = new_ws.cell(row_idx, col_idx)
                        if cell.value:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                    adjusted_width = min(max_length + 2, 50)
                    new_ws.column_dimensions[col_letter].width = adjusted_width
        else:
            new_wb = openpyxl.Workbook()
            new_ws = new_wb.active
            new_ws.title = 'Sheet1'
            
            for col_idx, header in enumerate(csv_headers, start=1):
                new_ws.cell(1, col_idx, value=header)
            
            for row_num, data_row in enumerate(rows, start=2):
                for col_idx, header in enumerate(csv_headers, start=1):
                    value = data_row.get(header, '')
                    converted = convert_value(value) if value else None
                    
                    # Convert duration 119 to 120
                    if header == 'duration3' and converted == 119:
                        converted = 120
                    
                    new_ws.cell(row_num, col_idx, value=converted)
            
            for col_idx in range(1, len(csv_headers) + 1):
                col_letter = get_column_letter(col_idx)
                max_length = 0
                for row_idx in range(1, new_ws.max_row + 1):
                    cell = new_ws.cell(row_idx, col_idx)
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                adjusted_width = min(max_length + 2, 50)
                new_ws.column_dimensions[col_letter].width = adjusted_width
        
        safe_filename = booking_num.replace('/', '-').replace('\\', '-').replace(' ', '_')
        
        if has_dallas:
            safe_filename = f"TAC_{safe_filename}"
        else:
            safe_filename = f"CTV_{safe_filename}"
        
        output_file = output_path / f"{safe_filename}.xlsx"
        new_wb.save(output_file)
        
        # Track if this file has unplaced spots
        if has_unplaced:
            unplaced_count = sum(1 for row in rows 
                if 'unplaced' in row.get('dateschedule', '').strip().lower() 
                or 'unplaced' in row.get('airtimep', '').strip().lower())
            files_with_unplaced.append((output_file.name, unplaced_count))
        
        market_list = ', '.join(sorted(markets_in_booking))
        print(f"  Created: {output_file.name} ({len(rows)} spots) - {template_name} - Markets: {market_list}")
    
    print(f"\n✓ Successfully created {len(bookings_data)} files in '{output_folder}' folder")
    
    # Display warning for files with unplaced spots
    if files_with_unplaced:
        print("\n" + "!" * 60)
        print("⚠ WARNING: UNPLACED SPOTS DETECTED")
        print("!" * 60)
        print(f"\nThe following {len(files_with_unplaced)} file(s) contain unplaced spots:")
        print("(These spots need air times assigned)\n")
        for filename, count in files_with_unplaced:
            print(f"  • {filename} - {count} unplaced spot(s)")
        print("\n" + "!" * 60)
    
    if ctv_template_wb:
        ctv_template_wb.close()
    if tac_template_wb:
        tac_template_wb.close()


if __name__ == "__main__":
    script_dir = Path(__file__).parent
    input_folder = script_dir / "input"
    output_folder = script_dir / "output"
    
    input_folder.mkdir(exist_ok=True)
    output_folder.mkdir(exist_ok=True)
    
    csv_files = list(input_folder.glob("*.csv"))
    
    if not csv_files:
        print("No CSV files found in the 'input' folder.")
        print(f"Please place your post-log CSV file in: {input_folder}")
        input("\nPress Enter to exit...")
        sys.exit(1)
    
    print("=" * 60)
    print("POST-LOG REPORT SORTER")
    print("=" * 60)
    
    # Ask which log type
    print("\nAre you pulling postlogs or prelogs?")
    print("  1. Postlogs (what has already aired)")
    print("  2. Prelogs (what is scheduled to air)")
    print()
    
    while True:
        try:
            log_choice = input("Enter 1 or 2: ").strip()
            if log_choice == "1":
                log_type = "Post"
                ctv_template_name = "CTVPostTemplate.xlsx"
                tac_template_name = "TACPostTemplate.xlsx"
                break
            elif log_choice == "2":
                log_type = "Pre"
                ctv_template_name = "CTVPreTemplate.xlsx"
                tac_template_name = "TACPreTemplate.xlsx"
                break
            else:
                print("Please enter 1 or 2")
        except KeyboardInterrupt:
            print("\n\nCancelled by user.")
            sys.exit(0)
    
    print(f"\nLog Type: {log_type}logs")
    
    ctv_template_file = script_dir / ctv_template_name
    tac_template_file = script_dir / tac_template_name
    
    ctv_exists = ctv_template_file.exists()
    tac_exists = tac_template_file.exists()
    
    print("\nTemplate Status:")
    if ctv_exists:
        print(f"  ✓ Crossings TV: {ctv_template_name}")
    else:
        print(f"  ✗ Crossings TV: {ctv_template_name} (NOT FOUND)")
    
    if tac_exists:
        print(f"  ✓ The Asian Channel: {tac_template_name}")
    else:
        print(f"  ✗ The Asian Channel: {tac_template_name} (NOT FOUND)")
    
    if not ctv_exists and not tac_exists:
        print("\n⚠ ERROR: No templates found!")
        print(f"Please create {ctv_template_name} and/or {tac_template_name}")
        input("\nPress Enter to exit...")
        sys.exit(1)
    
    print("\nNote: The script will automatically select the correct template")
    print("      based on the market (Dallas = TAC, Others = CTV)")
    
    print("\nAvailable CSV files in 'input' folder:\n")
    
    for idx, file in enumerate(csv_files, 1):
        file_size = file.stat().st_size / 1024
        print(f"  {idx}. {file.name} ({file_size:.1f} KB)")
    
    print("\n" + "=" * 60)
    
    while True:
        try:
            choice = input(f"\nEnter the number of the file to process (1-{len(csv_files)}): ").strip()
            choice_num = int(choice)
            if 1 <= choice_num <= len(csv_files):
                selected_file = csv_files[choice_num - 1]
                break
            else:
                print(f"Please enter a number between 1 and {len(csv_files)}")
        except ValueError:
            print("Please enter a valid number")
        except KeyboardInterrupt:
            print("\n\nCancelled by user.")
            sys.exit(0)
    
    print(f"\nProcessing: {selected_file.name}")
    print("=" * 60 + "\n")
    
    try:
        split_csv_by_booking(
            str(selected_file), 
            str(output_folder),
            str(ctv_template_file) if ctv_exists else None,
            str(tac_template_file) if tac_exists else None,
            'COD_CONTRATTO1'
        )
        print("\n" + "=" * 60)
        print("COMPLETE!")
        print("=" * 60)
        input("\nPress Enter to exit...")
    except Exception as e:
        print(f"\nError: {str(e)}")
        import traceback
        traceback.print_exc()
        input("\nPress Enter to exit...")
        sys.exit(1)
